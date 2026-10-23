import json
from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from flask import Flask, abort, jsonify, redirect, render_template, request, url_for
from werkzeug.exceptions import RequestEntityTooLarge

from actions import approve, create_input, propose_action, reject
from db import get_conn, init_db
from logic import analyze, dedup, load_postings
from verify import check_url

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024  # 5 MB upload limit
init_db()

ALL_POSTINGS = load_postings("data/postings.json")
POSTINGS = {p["id"]: p for p in ALL_POSTINGS}

MESSAGES = {
    "executed": "Approved. The resources were saved once.",
    "skipped": "Already handled earlier, so nothing was saved again.",
    "rejected": "Rejected. Nothing was saved.",
    "failed": "The save failed. Nothing was marked as saved.",
    "not_found": "That action was not found.",
}


def sample_resume():
    with open("data/resume_excel_only.txt", encoding="utf-8") as f:
        return f.read()


def extract_resume_file(uploaded_file):
    """Extract readable text from a PDF, DOCX, or TXT resume upload."""
    filename = (uploaded_file.filename or "").strip()
    suffix = Path(filename).suffix.lower()
    raw = uploaded_file.read()
    if not raw:
        raise ValueError("The uploaded file is empty.")

    if suffix == ".pdf":
        try:
            with fitz.open(stream=raw, filetype="pdf") as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        except Exception as exc:
            raise ValueError("Could not read this PDF. Please upload a valid, non-password-protected PDF.") from exc
    elif suffix == ".docx":
        try:
            doc = Document(BytesIO(raw))
            parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)
        except Exception as exc:
            raise ValueError("Could not read this DOCX. Please upload a valid Word document.") from exc
    elif suffix == ".txt":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Could not read this TXT file. Save it as UTF-8 and try again.") from exc
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the uploaded file. If it is a scanned PDF, use a text-based PDF or paste the resume text.")
    return text


def validate(form):
    errors = []
    resume = form.get("resume", "").strip()
    role = form.get("target_role", "").strip()
    city = form.get("city", "").strip()
    hours_raw = str(form.get("hours_per_week", "")).strip()

    if not resume:
        errors.append("Resume text is required. Upload a PDF, DOCX, or TXT file, or paste your resume text.")
    if not role:
        errors.append("Target role is required.")
    if not city:
        errors.append("City is required.")

    hours = None
    try:
        hours = int(hours_raw)
        if not 1 <= hours <= 80:
            errors.append("Hours per week must be between 1 and 80.")
    except (ValueError, TypeError):
        errors.append("Hours per week must be a whole number.")
    return errors, resume, role, city, hours


def save_results(input_id, unique, result):
    """Store the postings, counts and gaps so every result is traceable."""
    conn = get_conn()
    for p in unique:
        conn.execute(
            "INSERT OR REPLACE INTO postings "
            "(id, input_id, title, company, city, source, url, description, dedup_key) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (p["id"], input_id, p["title"], p["company"], p["city"],
             p["source"], p["url"], p["description"], p["dedup_key"]),
        )
    for skill, info in result["skill_counts"].items():
        conn.execute(
            "INSERT INTO skill_counts (input_id, skill, posting_count, posting_ids) "
            "VALUES (?, ?, ?, ?)",
            (input_id, skill, info["count"], ",".join(info["posting_ids"])),
        )
    for g in result["gaps"]:
        conn.execute(
            "INSERT INTO gaps (input_id, skill, status, resume_evidence, reason) "
            "VALUES (?, ?, ?, ?, ?)",
            (input_id, g["skill"], g["status"], g["evidence"], g["reason"]),
        )
    conn.commit()
    conn.close()


def latest_check(url):
    """Most recent verification row for a URL, or None if never checked."""
    conn = get_conn()
    row = conn.execute(
        "SELECT * FROM verifications WHERE url = ? ORDER BY id DESC LIMIT 1",
        (url,),
    ).fetchone()
    conn.close()
    return row


def cited_posting_ids(input_id):
    conn = get_conn()
    rows = conn.execute(
        "SELECT posting_ids FROM skill_counts WHERE input_id = ?", (input_id,)
    ).fetchall()
    conn.close()
    ids = set()
    for r in rows:
        ids.update(i for i in r["posting_ids"].split(",") if i)
    return sorted(ids)


def get_action(action_id):
    conn = get_conn()
    row = conn.execute(
        "SELECT * FROM actions WHERE action_id = ?", (action_id,)
    ).fetchone()
    conn.close()
    return row


def build_targets(row):
    """Everything this action verifies, in a fixed order: resources, then postings."""
    change = json.loads(row["approved_version"] or row["proposed_change"])
    targets = []
    for item in change["items"]:
        targets.append({
            "kind": "resource",
            "label": f"{item['skill']}: {item['title']}",
            "url": item["url"],
        })
    for pid in cited_posting_ids(row["input_id"]):
        p = POSTINGS[pid]
        targets.append({
            "kind": "posting",
            "label": f"{pid} {p['title']} at {p['company']}",
            "url": p["url"],
        })
    return targets


@app.errorhandler(RequestEntityTooLarge)
def handle_file_too_large(_error):
    form = {
        "resume": request.form.get("resume", ""),
        "target_role": request.form.get("target_role", ""),
        "city": request.form.get("city", ""),
        "hours_per_week": request.form.get("hours_per_week", ""),
    }
    return render_template(
        "index.html",
        form=form,
        errors=["The uploaded file is too large. Please choose a file smaller than 5 MB."],
        result=None,
    ), 413


@app.route("/api/verify/<action_id>/<int:index>", methods=["POST"])
def api_verify(action_id, index):
    """Check one target of an approved action and return what was observed."""
    row = get_action(action_id)
    if row is None:
        return jsonify({"error": "action not found"}), 404
    if row["approval_status"] != "executed":
        return jsonify({"error": "action is not approved yet"}), 409

    targets = build_targets(row)
    if not 0 <= index < len(targets):
        return jsonify({"error": "no such item"}), 404

    target = targets[index]
    result = check_url(target["url"], target["kind"])
    saved = latest_check(target["url"])
    return jsonify({
        "index": index,
        "is_live": result["is_live"],
        "status_code": result["status_code"],
        "checked_at": saved["checked_at"] if saved else "",
    })


@app.route("/")
def index():
    form = {
        "resume": sample_resume(),
        "target_role": "Data Analyst",
        "city": "Hyderabad",
        "hours_per_week": 10,
    }
    return render_template("index.html", form=form, errors=[], result=None)


@app.route("/analyze", methods=["POST"])
def run_analysis():
    form = {
        "resume": request.form.get("resume", ""),
        "target_role": request.form.get("target_role", ""),
        "city": request.form.get("city", ""),
        "hours_per_week": request.form.get("hours_per_week", ""),
    }

    uploaded_file = request.files.get("resume_file")
    if uploaded_file and uploaded_file.filename:
        try:
            # Uploaded file takes priority over pasted text.
            form["resume"] = extract_resume_file(uploaded_file)
        except ValueError as exc:
            return render_template("index.html", form=form, errors=[str(exc)], result=None)

    errors, resume, role, city, hours = validate(form)
    if errors:
        return render_template("index.html", form=form, errors=errors, result=None)

    matching = [
        p for p in ALL_POSTINGS
        if role.lower() in p["title"].lower() and p["city"].lower() == city.lower()
    ]
    result = analyze(matching, resume, target_role=role, city=city)
    input_id = create_input(resume, role, city, hours)

    result["input_id"] = input_id
    result["matched_count"] = len(matching)
    if result["status"] == "ok":
        unique = dedup(matching)
        save_results(input_id, unique, result)
        result["proposal"] = propose_action(input_id, result["gaps"])

        gap_status = {g["skill"]: g["status"] for g in result["gaps"]}
        result["chart"] = [
            {"skill": s, "count": i["count"], "status": gap_status[s], "ids": i["posting_ids"]}
            for s, i in result["skill_counts"].items()
        ]
        result["posting_info"] = {
            p["id"]: {
                "title": p["title"],
                "company": p["company"],
                "source": p["source"],
                "url": f"/job/{p['id']}",
            }
            for p in unique
        }

    return render_template("index.html", form=form, errors=[], result=result)


@app.route("/approve/<action_id>", methods=["POST"])
def approve_route(action_id):
    outcome = approve(action_id)["result"]
    return redirect(url_for("action_page", action_id=action_id, msg=outcome))


@app.route("/reject/<action_id>", methods=["POST"])
def reject_route(action_id):
    outcome = reject(action_id)["result"]
    return redirect(url_for("action_page", action_id=action_id, msg=outcome))


@app.route("/verify/<action_id>", methods=["POST"])
def verify_route(action_id):
    row = get_action(action_id)
    if row is None:
        abort(404)
    if row["approval_status"] == "executed":
        change = json.loads(row["approved_version"])
        for item in change["items"]:
            check_url(item["url"], "resource")
        for pid in cited_posting_ids(row["input_id"]):
            check_url(POSTINGS[pid]["url"], "posting")
    return redirect(url_for("action_page", action_id=action_id, msg="verified"))


@app.route("/action/<action_id>")
def action_page(action_id):
    row = get_action(action_id)
    if row is None:
        abort(404)

    targets = build_targets(row)
    for t in targets:
        t["check"] = latest_check(t["url"])

    change = json.loads(row["approved_version"] or row["proposed_change"])
    return render_template(
        "action.html",
        action_id=action_id,
        status=row["approval_status"],
        destination=change["destination"],
        targets=targets,
        message=MESSAGES.get(request.args.get("msg", ""), ""),
    )


@app.route("/job/<job_id>")
def job(job_id):
    p = POSTINGS.get(job_id)
    if p is None:
        abort(404)
    return (
        f"<h1>{p['title']} at {p['company']}</h1>"
        f"<p>{p['city']} | via {p['source']}</p>"
        f"<p>{p['description']}</p>"
    )


if __name__ == "__main__":
    app.run(port=5000, threaded=True)